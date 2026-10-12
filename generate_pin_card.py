# -*- coding: utf-8 -*-
"""生成《运动指令盒接线速查卡》(单页 A4)"""
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

doc = Document()

normal = doc.styles['Normal']
normal.font.name = '宋体'
normal.font.size = Pt(10.5)
normal.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')

def set_run(run, ea='宋体', ascii_font='宋体', size=10.5, bold=False, color=None):
    run.font.name = ascii_font
    run.font.size = Pt(size)
    run.font.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn('w:rFonts'))
    if rfonts is None:
        rfonts = OxmlElement('w:rFonts')
        rpr.append(rfonts)
    rfonts.set(qn('w:eastAsia'), ea)

def shade_cell(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:fill'), fill)
    tcPr.append(shd)

def title_run(text, size, bold=False, color=None):
    p = doc.add_paragraph()
    r = p.add_run(text)
    set_run(r, size=size, bold=bold, color=color)
    return p

def section(text):
    p = doc.add_paragraph()
    p.space_before = Pt(6)
    p.space_after = Pt(2)
    r = p.add_run(text)
    set_run(r, size=11, bold=True, color=(0x1F, 0x3B, 0x73))
    return p

def table(headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = 'Table Grid'
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = t.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = ''
        r = hdr[i].paragraphs[0].add_run(h)
        set_run(r, bold=True, size=9.5)
        shade_cell(hdr[i], 'DEEAF6')
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ''
            r = cells[i].paragraphs[0].add_run(str(v))
            set_run(r, size=9)
    if widths:
        for i, w in enumerate(widths):
            for row in t.rows:
                row.cells[i].width = Cm(w)
    return t

# 标题
tp = doc.add_paragraph()
tp.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = tp.add_run('运动指令盒接线速查卡')
set_run(r, size=18, bold=True, color=(0x1F, 0x3B, 0x73))
sp = doc.add_paragraph()
sp.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = sp.add_run('西门子 S7-200 SMART ST60 · 两轴（X/Y）步进电机 · 8 个指令盒（4 类 × 2 轴）')
set_run(r, size=10.5, color=(0x40, 0x40, 0x40))
sp2 = doc.add_paragraph()
sp2.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = sp2.add_run('EN 一律接 SM0.0（常通）；拖盒前先跑完运动向导生成 AXISx_* 符号')
set_run(r, size=8.5, color=(0xB0, 0x6A, 0x00))

W = [2.4, 2.0, 5.6, 5.6]

section('1. AXISx_CTRL（轴使能，每扫描周期必调）')
table(['引脚', '类型', 'AXIS0（X轴）', 'AXIS1（Y轴）'], [
    ['EN', '输入', 'SM0.0', 'SM0.0'],
    ['MOD_EN', '输入', 'M0.7', 'M0.7'],
    ['Done', '输出', 'M1.0', 'M1.2'],
    ['Error', '输出', 'VB100', 'VB102'],
    ['C_Pos', '输出', 'VD104', 'VD204'],
    ['C_Speed', '输出', 'VD108', 'VD208'],
    ['C_Dir', '输出', 'M1.1', 'M1.3'],
], widths=W)

section('2. AXISx_RSEEK（回原点）')
table(['引脚', '类型', 'AXIS0（X轴）', 'AXIS1（Y轴）'], [
    ['EN', '输入', 'SM0.0', 'SM0.0'],
    ['START', '输入', 'M2.1', 'M2.2'],
    ['Done', '输出', 'M4.0', 'M4.1'],
    ['Error', '输出', 'VB110', 'VB120'],
], widths=W)

section('3. AXISx_MAN（手动点动 / 连续）')
table(['引脚', '类型', 'AXIS0（X轴）', 'AXIS1（Y轴）'], [
    ['EN', '输入', 'SM0.0', 'SM0.0'],
    ['RUN', '输入', 'M0.3（0）', 'M0.3（0）'],
    ['JOG_P', '输入', 'I0.6', 'I1.0'],
    ['JOG_N', '输入', 'I0.7', 'I1.1'],
    ['Speed', '输入', 'VD130', 'VD230'],
    ['Dir', '输入', 'M0.4（0）', 'M0.4（0）'],
    ['Error', '输出', 'VB140', 'VB240'],
    ['C_Pos', '输出', 'VD144', 'VD244'],
    ['C_Speed', '输出', 'VD148', 'VD248'],
    ['C_Dir', '输出', 'M0.5', 'M0.6'],
], widths=W)

section('4. AXISx_GOTO（定位：Mode=0绝对 / 1相对）')
table(['引脚', '类型', 'AXIS0（X轴）', 'AXIS1（Y轴）'], [
    ['EN', '输入', 'SM0.0', 'SM0.0'],
    ['START', '输入', 'M3.0', 'M3.3'],
    ['Pos', '输入', 'VD300', 'VD400'],
    ['Speed', '输入', 'VD310', 'VD410'],
    ['Mode', '输入', '0', '0'],
    ['Abort', '输入', 'M3.1', 'M3.4'],
    ['Done', '输出', 'M3.2', 'M3.5'],
    ['Error', '输出', 'VB150', 'VB250'],
    ['C_Pos', '输出', 'VD154', 'VD254'],
    ['C_Speed', '输出', 'VD158', 'VD258'],
], widths=W)

# 备注
note = doc.add_paragraph()
note.space_before = Pt(6)
r = note.add_run('备注：表中「0」表示直接接常 0；Y 轴 GOTO 的地址（M3.3/M3.4/M3.5、VB250、VD254/VD258）为本卡新增分配。')
set_run(r, size=8.5, color=(0x60, 0x60, 0x60))

out = r'e:\两轴步进电机控制\运动指令盒接线速查卡.docx'
doc.save(out)
print('已生成:', out)